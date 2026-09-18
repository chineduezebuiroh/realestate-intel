import io
from datetime import datetime
from decimal import Decimal

import pytest

from scripts.nrc_b_verify import (
    COMPLETIONS, GEOGRAPHIES, METRICS, ProviderContractError, STARTS,
    canonical_hash, census_response_diagnostic, compare, main, month_end,
    geography_reconciliation,
    parse_census_response, parse_census_workbook, parse_fred_csv, parse_number,
    validate_rows, _provider_summary,
)


def row(geo="us_nation", metric=STARTS, period="2026-01", value="1500"):
    return {"geo_id": geo, "metric_id": metric, "date": month_end(period),
            "property_type_id": "all", "property_type": "all", "value": value,
            "provider": "fixture", "native_id": "x"}


def workbook(kind="starts", *, saar=True, units=True, headers=True,
             duplicate=False, unavailable=False, month_header=True,
             date_value=None, sheet_names=None, leading_month=None,
             footer_rows=None, all_missing=False, unavailable_token="(X)"):
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    names = sheet_names or [
        "Annual", "Not Seasonally Adjusted", "Seasonally Adjusted", "Seasonal Factors"]
    for name in names:
        book.create_sheet(name)
    if "Seasonally Adjusted" not in book.sheetnames:
        output = io.BytesIO(); book.save(output); book.close(); return output.getvalue()
    sheet = book["Seasonally Adjusted"]
    title = ("New Privately-Owned Housing Units Started" if kind == "starts" else
             "New Privately-Owned Housing Units Completed")
    sheet.append([title])
    sheet.append(["Seasonally adjusted annual rate" if saar else "Monthly values"])
    sheet.append(["Thousands of units. Detail may not add to total because of rounding."
                  if units else "Individual units"])
    sheet.append([])
    sheet.append(["Month" if month_header else "Period", "United States",
                  "Northeast", "Midwest", "South", "West"])
    sheet.append([None, "Total", "Total", "Total", "Total", "Total"])
    if leading_month is not None:
        sheet.append([leading_month])
    first_date = datetime(2026, 1, 1) if date_value is None else date_value
    first_values = (["(X)"] * 5 if all_missing else
                    [1500, unavailable_token if unavailable else 100, 200, 700, 500])
    sheet.append([first_date, *first_values])
    sheet.append([datetime(2026, 2, 1), 1501, 101, 201, 701, 498])
    if duplicate:
        sheet.append([first_date, 1500, 100, 200, 700, 500])
    if not headers:
        sheet.cell(6, 2).value = "All units"
    for footer in footer_rows or []:
        sheet.append(footer)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_numeric_missing_and_no_rescaling():
    assert parse_number("1,500") == Decimal("1500")
    assert parse_number(".") is None
    assert parse_number("(X)") is None
    assert row(value=str(parse_number("1500")))["value"] == "1500"
    with pytest.raises(ValueError, match="nonnumeric"): parse_number("secret")


@pytest.mark.parametrize(("source", "expected"), [
    ("2024-02", "2024-02-29"), ("2025-02-01", "2025-02-28"),
    ("January 2026", "2026-01-31")])
def test_month_end(source, expected): assert month_end(source) == expected


def test_workbook_parser_selects_five_totals_and_metric_identity():
    starts, contract = parse_census_workbook(workbook("starts"), "starts")
    completions, _ = parse_census_workbook(workbook("completions"), "completions")
    assert len(starts) == len(completions) == 10
    assert {x["geo_id"] for x in starts} == set(GEOGRAPHIES.values())
    assert {x["metric_id"] for x in starts} == {STARTS}
    assert {x["metric_id"] for x in completions} == {COMPLETIONS}
    assert starts[0]["value"] != "1500000"
    assert {x["date"] for x in starts} == {"2026-01-31", "2026-02-28"}
    assert contract["unit"] == "thousands_of_housing_units_saar"


def test_workbook_unavailable_marker_is_not_synthesized():
    rows, contract = parse_census_workbook(workbook(unavailable=True), "starts")
    assert len(rows) == 9
    assert contract["unavailable_cell_count_by_geography"] == {"Northeast": 1}


def test_completions_parenthesized_na_is_omitted_and_inventoried():
    assert parse_number("(NA)") is None
    rows, contract = parse_census_workbook(
        workbook("completions", unavailable=True, unavailable_token="(NA)"), "completions")
    assert len(rows) == 9
    assert contract["unavailable_cell_count_by_geography"] == {"Northeast": 1}
    with pytest.raises(ValueError, match="nonnumeric"):
        parse_number("(UNKNOWN)")


@pytest.mark.parametrize(("kwargs", "message"), [
    ({"saar": False}, "seasonally adjusted"), ({"units": False}, "thousands"),
    ({"headers": False}, "header"), ({"kind": "completions"}, "identity")])
def test_workbook_contract_drift_fails_closed(kwargs, message):
    requested_kind = "starts"
    with pytest.raises(ProviderContractError, match=message):
        parse_census_workbook(workbook(**kwargs), requested_kind)


def test_workbook_duplicate_keys_fail_closed():
    with pytest.raises(ValueError, match="duplicate"):
        parse_census_workbook(workbook(duplicate=True), "starts")


def test_real_month_header_is_required():
    with pytest.raises(ProviderContractError, match="header"):
        parse_census_workbook(workbook(month_header=False), "starts")


@pytest.mark.parametrize("bad_date", ["2026-01-01", 46023])
def test_month_date_must_deserialize_as_date_or_datetime(bad_date):
    with pytest.raises(ProviderContractError, match="non-date Month cell"):
        parse_census_workbook(workbook(date_value=bad_date), "starts")


def test_month_date_must_be_first_day():
    with pytest.raises(ProviderContractError, match="first-of-month"):
        parse_census_workbook(workbook(date_value=datetime(2026, 1, 2)), "starts")


def test_exact_sheet_inventory_is_required():
    with pytest.raises(ProviderContractError, match="sheets"):
        parse_census_workbook(workbook(sheet_names=["Annual", "Seasonally Adjusted"]), "starts")


@pytest.mark.parametrize("footer", [["Footnotes:"], ["arbitrary provider footer"]])
def test_non_date_footer_terminates_observation_block(footer):
    rows, _ = parse_census_workbook(workbook(footer_rows=[footer]), "starts")
    assert len(rows) == 10


def test_blank_then_footer_terminates_observation_block():
    rows, _ = parse_census_workbook(
        workbook(footer_rows=[[None], ["notes after blank row"]]), "starts")
    assert len(rows) == 10


def test_date_after_footer_fails_contiguous_block_contract():
    with pytest.raises(ProviderContractError, match="not one contiguous block"):
        parse_census_workbook(workbook(footer_rows=[
            ["notes"], [datetime(2026, 3, 1), 1502, 102, 202, 702, 496]]), "starts")


def test_missing_metrics_on_dated_row_do_not_terminate_block():
    rows, contract = parse_census_workbook(workbook(all_missing=True), "starts")
    assert len(rows) == 5
    assert contract["unavailable_cell_count_by_geography"] == {
        "Midwest": 1, "Northeast": 1, "South": 1, "United States": 1, "West": 1}


def test_non_date_before_first_observation_fails_closed():
    with pytest.raises(ProviderContractError, match="begins with a non-date"):
        parse_census_workbook(workbook(leading_month="unexpected structural row"), "starts")


def test_non_xlsx_diagnostic_is_bounded():
    body = b"<html><title>Missing Key</title>" + b"x" * 500
    metadata = {"status": 200, "content_type": "text/html",
                "requested_url": "https://example.test/request", "final_url": "https://example.test/error"}
    diagnostic = census_response_diagnostic(body, metadata)
    assert diagnostic["response_byte_length"] == len(body)
    assert len(diagnostic["raw_sha256"]) == 64
    assert len(diagnostic["body_prefix"].encode()) <= 256
    with pytest.raises(ProviderContractError, match=r"non-XLSX.*status=200.*Missing Key"):
        parse_census_response(body, metadata, "starts")


def test_xlsx_bytes_with_unexpected_content_type_fail_closed():
    with pytest.raises(ProviderContractError, match="non-XLSX"):
        parse_census_response(workbook(), {"status": 200, "content_type": "text/plain"}, "starts")


def test_main_persists_and_hashes_raw_response_before_parse(tmp_path, monkeypatch):
    body = b"upstream error page"
    monkeypatch.setattr("scripts.nrc_b_verify._get", lambda url: (
        body, {"requested_url": url, "final_url": url, "status": 200, "content_type": "text/plain"}))
    workspace = tmp_path / "evidence"
    rc = main(["--legacy-db", "data/market_serving.duckdb", "--workspace", str(workspace), "--skip-fred"])
    assert rc == 2
    assert (workspace / "raw/census_starts.xlsx").read_bytes() == body


def test_fred_schema_and_missing():
    rows = parse_fred_csv("DATE,HOUST\n2026-01-01,1500\n2026-02-01,.\n", "HOUST")
    assert rows == [row(value="1500", period="2026-01") | {"provider": "fred", "native_id": "HOUST"}]
    with pytest.raises(ValueError, match="schema"): parse_fred_csv("date,value\n2026-01-01,1\n", "HOUST")


def test_hashes_deterministic_and_inventory_excludes_value():
    a = [row(value="1500"), row(metric=COMPLETIONS, value="1400")]
    changed = [dict(a[0], value="1501"), a[1]]
    assert canonical_hash(a) == canonical_hash(reversed(a))
    assert canonical_hash(a) != canonical_hash(changed)
    assert canonical_hash(a, True) == canonical_hash(changed, True)


def test_per_series_history_summary_reports_continuity_and_gap():
    continuous = [row(period="2026-01"), row(period="2026-02")]
    summary = _provider_summary(continuous)
    series = summary["series"][f"{STARTS}|us_nation"]
    assert series["observation_count"] == 2
    assert series["continuous_within_observed_bounds"] is True
    gap = _provider_summary([row(period="2026-01"), row(period="2026-03")])
    assert gap["series"][f"{STARTS}|us_nation"]["missing_periods_within_bounds"] == ["2026-02-28"]


def test_geography_reconciliation_uses_governed_manifest():
    result = geography_reconciliation(__import__("pathlib").Path("config/geo_manifest.csv"))
    assert [item["governed_canonical_geo_id"] for item in result] == [
        "us_nation", "us_region_northeast", "us_region_midwest",
        "us_region_south", "us_region_west"]
    assert result[1]["legacy_ingestion_geo_id"] == "northeast_region__region"


def test_duplicate_and_unexpected_identity_fail_closed():
    with pytest.raises(ValueError, match="duplicate"): validate_rows([row(), row()])
    with pytest.raises(ValueError, match="applicability"): validate_rows([row(geo="invented")])


def test_exact_only_and_difference_classification():
    left = [row(value="1500"), row(metric=COMPLETIONS, value="1400")]
    right = [row(value="1500"), row(metric=COMPLETIONS, value="1399")]
    result = compare(left, right, "CENSUS_ONLY", "FRED_ONLY", "VALUE_DIFFERENCE")
    assert result["exact_match_count"] == 1
    assert result["value_difference_count"] == 1
    assert result["maximum_absolute_difference"] == "1"
