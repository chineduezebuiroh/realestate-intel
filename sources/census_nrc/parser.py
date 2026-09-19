"""Frozen semantic parser for Census NRC historical XLSX workbooks."""
from __future__ import annotations

import calendar
import io
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

SOURCE_ID = "census_nrc"
STARTS = "census_housing_starts_total_saar"
COMPLETIONS = "census_housing_completions_total_saar"
METRICS = (STARTS, COMPLETIONS)
GEOGRAPHIES = {"US": "us_nation", "NE": "us_region_northeast",
               "MW": "us_region_midwest", "S": "us_region_south", "W": "us_region_west"}
CENSUS_INPUTS = {
    "starts": ("https://www.census.gov/construction/nrc/xls/starts_cust.xlsx", STARTS),
    "completions": ("https://www.census.gov/construction/nrc/xls/comps_cust.xlsx", COMPLETIONS),
}
KEY_FIELDS = ("geo_id", "metric_id", "date", "property_type_id")
PARSER_CONTRACT_VERSION = "census_nrc_workbook_parser_v1_openpyxl_3.1.5"


class ProviderContractError(ValueError):
    """Provider bytes do not satisfy the frozen workbook contract."""


def month_end(value: str | date | datetime) -> str:
    if isinstance(value, (date, datetime)):
        year, month = value.year, value.month
    else:
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%Y-%m", "%B %Y", "%b %Y"):
            try:
                parsed = datetime.strptime(text, fmt); year, month = parsed.year, parsed.month; break
            except ValueError: pass
        else: raise ValueError(f"invalid monthly period: {value!r}")
    return f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"


def parse_number(value: Any) -> Decimal | None:
    """Recognize only the provider's exact unavailable token (plus empty cells)."""
    text = "" if value is None else str(value).strip()
    if text in {"", "(NA)"}: return None
    try: number = Decimal(text.replace(",", ""))
    except InvalidOperation as exc: raise ValueError(f"unexpected nonnumeric value: {value!r}") from exc
    if not number.is_finite(): raise ValueError(f"non-finite value: {value!r}")
    return number


def _row(geo_code: str, metric: str, period: Any, value: Any, provider: str = "census",
         native_id: str = "") -> dict[str, Any] | None:
    if geo_code not in GEOGRAPHIES or metric not in METRICS:
        raise ValueError(f"unexpected NRC identity: {geo_code}/{metric}")
    number = parse_number(value)
    if number is None: return None
    return {"geo_id": GEOGRAPHIES[geo_code], "metric_id": metric, "date": month_end(period),
            "property_type_id": "all", "property_type": "all", "value": str(number),
            "provider": provider, "native_id": native_id}


def validate_rows(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    result = sorted(rows, key=lambda x: tuple(str(x[k]) for k in KEY_FIELDS)); seen = set()
    for row in result:
        key = tuple(str(row[k]) for k in KEY_FIELDS)
        if key in seen: raise ValueError(f"duplicate normalized key: {key}")
        seen.add(key)
    extra = {(r["geo_id"], r["metric_id"]) for r in result} - {
        (g, m) for g in GEOGRAPHIES.values() for m in METRICS}
    if extra: raise ValueError(f"unexpected applicability pairs: {sorted(extra)}")
    return result


def parse_census_workbook(payload: bytes, kind: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if kind not in CENSUS_INPUTS: raise ValueError(f"unexpected Census workbook kind: {kind}")
    metric = CENSUS_INPUTS[kind][1]
    try: workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError, BadZipFile) as exc:
        raise ProviderContractError(f"Census workbook could not be decoded as XLSX: {exc}") from exc
    try:
        sheets = workbook.sheetnames
        expected = ["Annual", "Not Seasonally Adjusted", "Seasonally Adjusted", "Seasonal Factors"]
        if sheets != expected: raise ProviderContractError(f"unexpected Census workbook sheets: {sheets}")
        matrix = [list(row) for row in workbook["Seasonally Adjusted"].iter_rows(values_only=True)]
    finally: workbook.close()
    text = "\n".join(str(v) for row in matrix[:30] for v in row if v is not None).lower()
    phrase = "housing units started" if kind == "starts" else "housing units completed"
    if phrase not in text: raise ProviderContractError(f"Census {kind} workbook title/metric identity not found")
    if "seasonally adjusted annual rate" not in text: raise ProviderContractError("Census workbook lacks seasonally adjusted annual rate declaration")
    if "thousands of units" not in text: raise ProviderContractError("Census workbook lacks thousands-of-units declaration")
    provider_geos = {"United States": "US", "Northeast": "NE", "Midwest": "MW", "South": "S", "West": "W"}
    header = None
    for index in range(len(matrix) - 1):
        upper, lower = matrix[index], matrix[index + 1]
        if not upper or str(upper[0]).strip() != "Month": continue
        headers = [str(v).strip() for v in upper[1:] if v not in (None, "")]
        if len(headers) != 5 or set(headers) != set(provider_geos):
            raise ProviderContractError(f"unexpected Census governed geography headers: {headers}")
        columns, current = {}, None
        for col in range(max(len(upper), len(lower))):
            top = str(upper[col]).strip() if col < len(upper) and upper[col] is not None else ""
            bottom = str(lower[col]).strip() if col < len(lower) and lower[col] is not None else ""
            if top in provider_geos: current = top
            if current and bottom == "Total":
                if current in columns: raise ProviderContractError(f"duplicate Census Total header: {current}")
                columns[current] = col
        if set(columns) == set(provider_geos): header = (index + 1, columns); break
    if header is None: raise ProviderContractError("Census workbook two-row geography/Total header not found")
    header_row, columns = header; rows = []; unavailable = Counter(); started = terminated = False
    for source in matrix[header_row + 1:]:
        observed = source[0] if source else None
        if not isinstance(observed, (date, datetime)):
            if not started:
                if observed in (None, ""): continue
                raise ProviderContractError(f"Census monthly observation block begins with a non-date Month cell: {observed!r}")
            terminated = True; continue
        if terminated: raise ProviderContractError(f"Census monthly observations are not one contiguous block; found date after footer: {observed!r}")
        started = True
        if observed.day != 1: raise ProviderContractError(f"Census Month value is not first-of-month: {observed.isoformat()}")
        for label, col in columns.items():
            raw = source[col] if col < len(source) else None
            if parse_number(raw) is None: unavailable[label] += 1; continue
            rows.append(_row(provider_geos[label], metric, observed, raw, native_id=kind))
    normalized = validate_rows(r for r in rows if r is not None)
    if {(r["geo_id"], r["metric_id"]) for r in normalized} != {(g, metric) for g in GEOGRAPHIES.values()}:
        raise ProviderContractError(f"Census {kind} workbook lacks one or more governed series")
    return normalized, {"workbook_kind": kind, "sheet_names": sheets, "worksheet": "Seasonally Adjusted",
        "date_cell_contract": "openpyxl date/datetime; first day of month",
        "unit": "thousands_of_housing_units_saar", "numeric_scale_factor": 1,
        "parser_contract_version": PARSER_CONTRACT_VERSION,
        "unavailable_cell_count_by_geography": dict(sorted(unavailable.items()))}


def canonical_hash(rows: Iterable[dict[str, Any]], inventory_only: bool = False) -> str:
    import hashlib
    material = []
    for row in validate_rows(rows):
        fields = [str(row[k]) for k in KEY_FIELDS]
        if not inventory_only: fields.append(str(Decimal(str(row["value"]))))
        material.append("\x1f".join(fields))
    return hashlib.sha256(("\n".join(material) + "\n").encode()).hexdigest()
